import openpyxl
from tabulate import tabulate

excel_dataframe = openpyxl.load_workbook('personas.xlsx')
dataframe = excel_dataframe.active

data = []

for row in range(1, dataframe.max_row):
    _row = [row]

    for col in dataframe.iter_cols(1, dataframe.max_row):
        _row.append(col[row].value)

    data.append(_row)

headers = ["#", "id", "name", "company", "Email", "Mac Address"]
headers_align = ("center",) * 6

print(tabulate(data, headers=headers, tablefmt="fancy_grid", colalign=headers_align))

### Ver todos
def view_all():
    print(tabulate(data, headers=headers, tablefmt="fancy_grid", colalign=headers_align))


### 1 Name y Email
def name_email():
    lista_name_email = []

    for persona in data:
        lista_name_email.append([persona[2], persona[4]])

    print(tabulate(lista_name_email, headers=["Name", "Email"], tablefmt="fancy_grid"))


### 2 CONTAR REGISTROS
def count_records():
    print(f"\nTotal de registros: {len(data)} registros.")


### 3 empresas que contengan "Group"
def group_companies():
    lista_group = []

    for persona in data:
        if "Group" in persona[3]:
            lista_group.append(persona)

    print(tabulate(lista_group, headers=headers, tablefmt="fancy_grid", colalign=headers_align))


### 4 buscar una persona por su ID
def search_person_by_id():
    search_id = input("\nIngrese el ID de la persona que desea buscar: ").strip()
    search_result = [persona for persona in data if persona[1] == search_id]

    if search_result:
        print(tabulate(search_result, headers=headers, tablefmt="fancy_grid", colalign=headers_align))
    else:
        print("Persona no encontrada.")


#### 5 menu interactivo
while True:
    print("\n===== MENÚ =====")
    print("1. Ver todos")
    print("2. Ver solo Name y Email")
    print("3. Buscar por ID")
    print("4. Filtrar empresa")
    print("5. Contar registros")
    print("6. Salir")

    option = input("Seleccione una opción: ").strip()

    if option == "1":
        view_all()
    elif option == "2":
        name_email()
    elif option == "3":
        search_person_by_id()
    elif option == "4":
        group_companies()
    elif option == "5":
        count_records()
    elif option == "6":
        print("Hasta pronto")
        break
    else:
        print("Opción no válida, escriba un número del 1 al 6")
